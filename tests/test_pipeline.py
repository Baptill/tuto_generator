"""Tests de bout en bout du pipeline Étape 1 (voir CLAUDE.md, DoD Étape 1)."""
from __future__ import annotations

from docx import Document

from app import storage
from app.build import generate_article
from app.pagination import paginate
from app.schemas import Section
from app.validation import validate_article


def _section(id_: str, hauteur: int) -> Section:
    return Section(id=id_, titre=id_, hauteur=hauteur, blocs=[])


def test_pagination_respects_grid():
    sections = [_section("a", 1), _section("b", 1), _section("c", 2), _section("d", 4), _section("e", 3)]
    pages = paginate(sections)
    assert [[s.id for s in page] for page in pages] == [["a", "b", "c"], ["d"], ["e"]]
    for page in pages:
        assert sum(s.hauteur for s in page) <= 4


def test_validate_article_rejects_too_long_title():
    article, _ = storage.load_article("2026-07-exemple-tuto")
    config = storage.load_template_config(article.template_id)
    article.titre = "x" * 200
    report = validate_article(article, config, assets_dir=storage.article_dir(article.id))
    assert not report.ok
    assert "Titre trop long" in report.errors[0]


def test_generate_sample_article_end_to_end():
    result = generate_article("2026-07-exemple-tuto", skip_pdf=True)
    assert result.docx_path.is_file()
    assert len(result.hash_contenu) == 64

    doc = Document(str(result.docx_path))
    assert len(doc.inline_shapes) == 3  # 3 images, cf. DoD Étape 1
    styles_used = {p.style.name for p in doc.paragraphs}
    assert {"Titre 1", "Titre 2", "Corps", "Legende", "Encadre Astuce", "Encadre Attention", "Encadre Info"} <= styles_used

    page_breaks = sum(
        1
        for p in doc.paragraphs
        for r in p.runs
        if 'w:type="page"' in r._element.xml
    )
    assert page_breaks == 2  # 3 pages : (1+1+2), (4), (3) — cf. content.yaml de l'article
