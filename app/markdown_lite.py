"""Markdown léger → runs Word (voir CLAUDE.md, Étape 1, « Points techniques délicats »).

Supporté volontairement minimal :
- **gras**
- lignes commençant par « - » ou « * » → liste à puces
- paragraphes séparés par des lignes vides

Tout le reste est traité comme du texte brut. Pas de markdown imbriqué, pas de
liens, pas de listes numérotées — à étendre si le besoin apparaît en pratique.
"""
from __future__ import annotations

import re
from typing import Iterable

from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_BULLET_RE = re.compile(r"^\s*[-*]\s+(.*)$")


def _add_runs_with_bold(paragraph: Paragraph, text: str) -> None:
    """Découpe `text` sur les marqueurs **gras** et ajoute les runs correspondants."""
    pos = 0
    for match in _BOLD_RE.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        bold_run = paragraph.add_run(match.group(1))
        bold_run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _split_blocks(texte: str) -> Iterable[tuple[str, list[str]]]:
    """Regroupe les lignes en blocs ('paragraphe' | 'liste', [lignes])."""
    lines = [line for line in texte.strip("\n").splitlines()]
    current_kind: str | None = None
    buffer: list[str] = []

    def flush():
        if buffer:
            yield_kind = current_kind or "paragraphe"
            return yield_kind, list(buffer)
        return None

    blocks: list[tuple[str, list[str]]] = []
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            if buffer:
                blocks.append((current_kind or "paragraphe", buffer))
                buffer = []
                current_kind = None
            continue
        bullet_match = _BULLET_RE.match(line)
        kind = "liste" if bullet_match else "paragraphe"
        content = bullet_match.group(1) if bullet_match else line
        if buffer and kind != current_kind:
            blocks.append((current_kind or "paragraphe", buffer))
            buffer = []
        current_kind = kind
        buffer.append(content)
    if buffer:
        blocks.append((current_kind or "paragraphe", buffer))
    return blocks


def add_markdown_lite(
    container: DocxDocument,
    texte: str,
    style: str | None = None,
    list_style: str = "List Bullet",
) -> None:
    """Ajoute `texte` (markdown léger) au conteneur docx (`Document` ou sous-document
    docxtpl), en appliquant `style` aux paragraphes normaux et `list_style` aux
    listes à puces.
    """
    for kind, lines in _split_blocks(texte):
        if kind == "liste":
            for line in lines:
                p = container.add_paragraph(style=list_style)
                _add_runs_with_bold(p, line)
        else:
            p = container.add_paragraph(style=style)
            _add_runs_with_bold(p, " ".join(lines))
