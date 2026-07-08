"""Catalogue des layouts de section.

Un layout décrit la disposition visuelle des blocs d'une section dans le Word.
Le champ `layout` du content.yaml choisit lequel appliquer :

  texte-seul        — titre puis blocs empilés verticalement (défaut)
  image-dessus-texte — image pleine largeur, texte en dessous
  image-gauche-texte — tableau 2 col : image | texte
  texte-image-droite — tableau 2 col : texte | image

Chaque fonction reçoit les mêmes arguments et retourne un DocxSubdoc prêt à
être injecté dans le template via la boucle `{%p for el in elements %}`.
"""
from __future__ import annotations

from pathlib import Path

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm
from docxtpl import DocxTemplate

from app.markdown_lite import add_markdown_lite
from app.schemas import BlocEncadre, BlocImage, BlocParagraphe, Section

_STYLE_TITRE_2 = "Titre 2"
_STYLE_CORPS = "Corps"
_STYLE_LEGENDE = "Legende"
_ENCADRE_STYLES = {
    "astuce": "Encadre Astuce",
    "attention": "Encadre Attention",
    "info": "Encadre Info",
}

# Largeurs des colonnes du layout 2 colonnes (A4 = 210mm − 2×20mm marges = 170mm)
_COL_IMAGE_MM = 75
_COL_TEXTE_MM = 90


# ---------------------------------------------------------------------------
# Helpers internes
# ---------------------------------------------------------------------------

def _remove_table_borders(table) -> None:
    """Rend le tableau invisible (pas de trait de bordure)."""
    tbl = table._tbl
    tblPr = tbl.find(qn("w:tblPr"))
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "none")
        tblBorders.append(el)
    tblPr.append(tblBorders)


def _set_col_width(table, col_idx: int, width_mm: int) -> None:
    for cell in table.columns[col_idx].cells:
        cell.width = Mm(width_mm)


def _add_image_bloc(container, bloc: BlocImage, article_dir: Path, largeur_mm: int) -> None:
    image_path = article_dir / bloc.fichier
    p = container.add_paragraph()
    p.add_run().add_picture(str(image_path), width=Mm(largeur_mm))
    if bloc.legende:
        container.add_paragraph(bloc.legende, style=_STYLE_LEGENDE)


def _add_texte_blocs(container, blocs, titre: str | None = None) -> None:
    if titre:
        container.add_paragraph(titre, style=_STYLE_TITRE_2)
    for bloc in blocs:
        if isinstance(bloc, BlocParagraphe):
            add_markdown_lite(container, bloc.texte, style=_STYLE_CORPS)
        elif isinstance(bloc, BlocEncadre):
            style = _ENCADRE_STYLES[bloc.style]
            add_markdown_lite(container, bloc.texte, style=style, list_style=style)


# ---------------------------------------------------------------------------
# Layouts
# ---------------------------------------------------------------------------

def layout_texte_seul(
    tpl: DocxTemplate, section: Section, article_dir: Path, largeur_mm_defaut: int
):
    """Titre puis blocs empilés verticalement dans l'ordre du content.yaml."""
    subdoc = tpl.new_subdoc()
    subdoc.add_paragraph(section.titre, style=_STYLE_TITRE_2)
    for bloc in section.blocs:
        if isinstance(bloc, BlocParagraphe):
            add_markdown_lite(subdoc, bloc.texte, style=_STYLE_CORPS)
        elif isinstance(bloc, BlocEncadre):
            style = _ENCADRE_STYLES[bloc.style]
            add_markdown_lite(subdoc, bloc.texte, style=style, list_style=style)
        elif isinstance(bloc, BlocImage):
            _add_image_bloc(subdoc, bloc, article_dir, bloc.largeur_mm or largeur_mm_defaut)
    return subdoc


def layout_image_dessus_texte(
    tpl: DocxTemplate, section: Section, article_dir: Path, largeur_mm_defaut: int
):
    """Titre, puis image(s) pleine largeur, puis blocs texte/encadrés."""
    subdoc = tpl.new_subdoc()
    subdoc.add_paragraph(section.titre, style=_STYLE_TITRE_2)

    images = [b for b in section.blocs if isinstance(b, BlocImage)]
    autres = [b for b in section.blocs if not isinstance(b, BlocImage)]

    for bloc in images:
        _add_image_bloc(subdoc, bloc, article_dir, bloc.largeur_mm or largeur_mm_defaut)
    _add_texte_blocs(subdoc, autres)

    return subdoc


def _layout_deux_colonnes(
    tpl: DocxTemplate,
    section: Section,
    article_dir: Path,
    largeur_mm_defaut: int,
    image_a_gauche: bool,
) -> object:
    """Tableau 2 colonnes sans bordure : image d'un côté, texte de l'autre."""
    subdoc = tpl.new_subdoc()

    images = [b for b in section.blocs if isinstance(b, BlocImage)]
    autres = [b for b in section.blocs if not isinstance(b, BlocImage)]

    table = subdoc.add_table(rows=1, cols=2)
    table.autofit = False
    _remove_table_borders(table)

    col_img = 0 if image_a_gauche else 1
    col_txt = 1 if image_a_gauche else 0

    _set_col_width(table, col_img, _COL_IMAGE_MM)
    _set_col_width(table, col_txt, _COL_TEXTE_MM)

    if images:
        _add_image_bloc(table.cell(0, col_img), images[0], article_dir, _COL_IMAGE_MM)

    _add_texte_blocs(table.cell(0, col_txt), autres, titre=section.titre)

    return subdoc


def layout_image_gauche_texte(
    tpl: DocxTemplate, section: Section, article_dir: Path, largeur_mm_defaut: int
):
    """Tableau 2 col : image à gauche, titre + texte à droite."""
    return _layout_deux_colonnes(tpl, section, article_dir, largeur_mm_defaut, image_a_gauche=True)


def layout_texte_image_droite(
    tpl: DocxTemplate, section: Section, article_dir: Path, largeur_mm_defaut: int
):
    """Tableau 2 col : titre + texte à gauche, image à droite."""
    return _layout_deux_colonnes(tpl, section, article_dir, largeur_mm_defaut, image_a_gauche=False)


# ---------------------------------------------------------------------------
# Dispatch
# ---------------------------------------------------------------------------

LAYOUTS = {
    "texte-seul": layout_texte_seul,
    "image-dessus-texte": layout_image_dessus_texte,
    "image-gauche-texte": layout_image_gauche_texte,
    "texte-image-droite": layout_texte_image_droite,
}


def build_section_subdoc(
    tpl: DocxTemplate, section: Section, article_dir: Path, largeur_mm_defaut: int
):
    fn = LAYOUTS.get(section.layout)
    if fn is None:
        raise ValueError(f"Layout inconnu : '{section.layout}'. Valeurs possibles : {list(LAYOUTS)}")
    return fn(tpl, section, article_dir, largeur_mm_defaut)
