"""Moteur de rendu — content.yaml + template.docx + charte.yaml → article.docx.

Voir CLAUDE.md §3 (sections modulaires, styles nommés) et §4 (modèle de
données). Règle d'or : ce module lit `content.yaml`, jamais l'inverse.
"""
from __future__ import annotations

from pathlib import Path

from docx.enum.text import WD_BREAK
from docx.shared import Mm, RGBColor
from docxtpl import DocxTemplate

from app.markdown_lite import add_markdown_lite
from app.pagination import paginate
from app.schemas import Article, BlocEncadre, BlocImage, BlocParagraphe, Charte, Section, TemplateConfig

STYLE_TITRE_1 = "Titre 1"
STYLE_TITRE_2 = "Titre 2"
STYLE_CORPS = "Corps"
STYLE_LEGENDE = "Legende"
ENCADRE_STYLES = {
    "astuce": "Encadre Astuce",
    "attention": "Encadre Attention",
    "info": "Encadre Info",
}


class RenderError(Exception):
    pass


def _build_section_subdoc(tpl: DocxTemplate, section: Section, article_dir: Path, largeur_mm_defaut: int):
    subdoc = tpl.new_subdoc()
    subdoc.add_paragraph(section.titre, style=STYLE_TITRE_2)

    for bloc in section.blocs:
        if isinstance(bloc, BlocParagraphe):
            add_markdown_lite(subdoc, bloc.texte, style=STYLE_CORPS)
        elif isinstance(bloc, BlocEncadre):
            style_name = ENCADRE_STYLES[bloc.style]
            add_markdown_lite(subdoc, bloc.texte, style=style_name, list_style=style_name)
        elif isinstance(bloc, BlocImage):
            image_path = article_dir / bloc.fichier
            if not image_path.is_file():
                raise RenderError(f"Section '{section.id}' : image introuvable '{image_path}'.")
            width_mm = bloc.largeur_mm or largeur_mm_defaut
            p = subdoc.add_paragraph()
            run = p.add_run()
            run.add_picture(str(image_path), width=Mm(width_mm))
            if bloc.legende:
                subdoc.add_paragraph(bloc.legende, style=STYLE_LEGENDE)
        else:  # pragma: no cover - garde-fou si le schéma évolue
            raise RenderError(f"Type de bloc non géré : {bloc!r}")

    return subdoc


def _build_page_break_subdoc(tpl: DocxTemplate):
    subdoc = tpl.new_subdoc()
    run = subdoc.add_paragraph().add_run()
    run.add_break(WD_BREAK.PAGE)
    return subdoc


def _apply_charte(tpl: DocxTemplate, charte: Charte) -> None:
    """Applique couleurs/polices/espacements de la charte aux styles nommés du
    document rendu (voir CLAUDE.md §3, « Template Word : styles nommés »).

    Les couleurs sémantiques des encadrés (astuce/attention/info) restent
    fixes dans le template : elles portent un sens indépendant de la marque,
    contrairement aux couleurs de titres/corps qui, elles, suivent la charte.
    """
    document = tpl.docx
    styles = document.styles
    color_primaire = RGBColor.from_string(charte.couleurs.primaire.lstrip("#").upper())
    color_texte = RGBColor.from_string(charte.couleurs.texte.lstrip("#").upper())

    style_names = {s.name for s in styles}

    for name in (STYLE_TITRE_1, STYLE_TITRE_2):
        if name in style_names:
            font = styles[name].font
            font.name = charte.polices.titres
            font.color.rgb = color_primaire

    if STYLE_CORPS in style_names:
        corps = styles[STYLE_CORPS]
        corps.font.name = charte.polices.corps
        corps.font.color.rgb = color_texte
        corps.paragraph_format.line_spacing = charte.espacements.interligne

    if STYLE_LEGENDE in style_names:
        styles[STYLE_LEGENDE].font.name = charte.polices.corps

    for section in document.sections:
        marge = Mm(charte.espacements.marge_mm)
        section.left_margin = marge
        section.right_margin = marge
        section.top_margin = marge
        section.bottom_margin = marge


def render_article(
    article: Article,
    config: TemplateConfig,
    charte: Charte,
    template_path: Path,
    article_dir: Path,
    output_path: Path,
) -> Path:
    """Rend `article` avec `template_path` + `charte`, écrit le docx dans
    `output_path`. `article_dir` est le dossier contenant `assets/` (les
    chemins d'image du content.yaml sont relatifs à ce dossier).
    """
    tpl = DocxTemplate(str(template_path))

    pages = paginate(article.sections)
    elements = []
    for i, page_sections in enumerate(pages):
        if i > 0:
            elements.append(_build_page_break_subdoc(tpl))
        for section in page_sections:
            elements.append(
                _build_section_subdoc(tpl, section, article_dir, config.image.largeur_mm_defaut)
            )

    context = {"article": article, "elements": elements}
    tpl.render(context)

    _apply_charte(tpl, charte)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    tpl.save(str(output_path))
    return output_path
