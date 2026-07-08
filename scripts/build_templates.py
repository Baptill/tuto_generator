"""Génère les templates Word pilotes `tuto-release` et `blog-standard`.

Voir CLAUDE.md, Étape 1, Lot 1.B « Templates pilotes » et §6 « Guide
d'authoring d'un template ». Ce script joue le rôle du designer qui
construirait le .docx dans Word : styles nommés, balises Jinja, pas de
couleur en dur (les couleurs de titres/corps sont injectées à chaque build
depuis charte.yaml — voir app/renderer.py::_apply_charte).

Usage : python -m scripts.build_templates
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

VAULT_ROOT = Path(__file__).resolve().parent.parent / "vault-articles"

# Couleurs sémantiques des encadrés : fixes, indépendantes de la charte de
# marque (voir app/renderer.py::_apply_charte pour la justification).
ENCADRE_FILLS = {
    "Encadre Astuce": "E8F5E9",  # vert clair
    "Encadre Attention": "FDECEA",  # rouge/orangé clair
    "Encadre Info": "E8F0FE",  # bleu clair
}
ENCADRE_BORDER_COLORS = {
    "Encadre Astuce": "2E7D32",
    "Encadre Attention": "C62828",
    "Encadre Info": "1565C0",
}


def _set_shading(paragraph_format_element, fill_hex: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    paragraph_format_element.append(shd)


def _set_left_border(paragraph_format_element, color_hex: str) -> None:
    pBdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "4")
    left.set(qn("w:color"), color_hex)
    pBdr.append(left)
    paragraph_format_element.append(pBdr)


def _add_style(doc: Document, name: str, base: str = "Normal", size: int | None = None, bold: bool = False):
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles[base]
    if size:
        style.font.size = Pt(size)
    style.font.bold = bold
    return style


def _build_named_styles(doc: Document) -> None:
    titre1 = _add_style(doc, "Titre 1", size=24, bold=True)
    titre1.paragraph_format.space_after = Pt(12)
    titre1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    titre2 = _add_style(doc, "Titre 2", size=16, bold=True)
    titre2.paragraph_format.space_before = Pt(16)
    titre2.paragraph_format.space_after = Pt(6)

    corps = _add_style(doc, "Corps", size=11)
    corps.paragraph_format.space_after = Pt(8)

    legende = _add_style(doc, "Legende", size=9)
    legende.font.italic = True
    legende.paragraph_format.space_after = Pt(10)
    legende.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name, fill in ENCADRE_FILLS.items():
        style = _add_style(doc, name, size=11)
        pPr = style.element.get_or_add_pPr()
        _set_shading(pPr, fill)
        _set_left_border(pPr, ENCADRE_BORDER_COLORS[name])
        style.paragraph_format.space_before = Pt(6)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.left_indent = Pt(10)


def _build_body(doc: Document) -> None:
    doc.add_paragraph("{{ article.titre }}", style="Titre 1")
    # Balises Jinja de docxtpl : `{%p %}` = tag « paragraphe » de contrôle
    # (for/endfor) et `{{p el }}` = tag « paragraphe » d'impression — les deux
    # variantes suppriment leur `<w:p>` englobant au moment du patch XML, ce
    # qui est indispensable ici : `el` est un sous-document (XML brut d'une
    # section ou d'un saut de page, voir app/renderer.py) et doit être inséré
    # tel quel dans le corps du document, pas encapsulé dans un `<w:t>`.
    doc.add_paragraph("{%p for el in elements %}")
    doc.add_paragraph("{{p el }}")
    doc.add_paragraph("{%p endfor %}")


def build_template(template_id: str, libelle: str) -> Path:
    doc = Document()
    _build_named_styles(doc)
    _build_body(doc)

    out_dir = VAULT_ROOT / "templates" / template_id
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "template.docx"
    doc.save(str(out_path))
    return out_path


def main() -> None:
    for template_id, libelle in [
        ("tuto-release", "Tutoriel de release"),
        ("blog-standard", "Article de blog"),
    ]:
        path = build_template(template_id, libelle)
        print(f"{template_id} -> {path}")


if __name__ == "__main__":
    main()
